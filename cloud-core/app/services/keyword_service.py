"""Knowledge Object keyword & query workspace (REQ-KW-001..004, ADR-092).
Keywords are knowledge inputs: they drive research tasks (source discovery),
and every created task is traced back to its keyword."""
from __future__ import annotations

import csv
import io
import json
import unicodedata
import uuid
from uuid import UUID

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.models.ai_center import KEYWORD_KINDS, KnowledgeKeyword
from app.models.enums import ResearchMode
from app.models.knowledge_object import KnowledgeObject
from app.services import research_project_service, research_task_service

MAX_IMPORT_BYTES = 2_000_000
MAX_IMPORT_ROWS = 5000
KEYWORD_COLUMNS = ("keyword", "query", "term", "کلمه کلیدی")


def normalize(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).strip().lower().split())


def _ko(db: Session, organization_id: UUID, ko_id: UUID) -> KnowledgeObject:
    ko = db.execute(select(KnowledgeObject).where(KnowledgeObject.id == ko_id,
                                                  KnowledgeObject.organization_id == organization_id)).scalar_one_or_none()
    if ko is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Knowledge object not found")
    return ko


def _int_or_none(v):  # noqa: ANN001, ANN202
    try:
        return int(str(v).replace(",", "").strip()) if v not in (None, "") else None
    except ValueError:
        return None


def rows_from_upload(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File larger than 2 MB")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception:  # noqa: BLE001
            raise HTTPException(status_code=422, detail="Not a readable .xlsx workbook") from None
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it, [])]
        return [{header[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(header)} for r in it]
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=422, detail="File must be UTF-8 text, CSV, JSON or .xlsx") from None
    if name.endswith(".json"):
        try:
            payload = json.loads(text)
        except ValueError:
            raise HTTPException(status_code=422, detail="Invalid JSON") from None
        if not isinstance(payload, list):
            raise HTTPException(status_code=422, detail="JSON must be a list of strings or objects")
        return [{"keyword": p} if isinstance(p, str) else {str(k).lower(): v for k, v in p.items()} for p in payload
                if isinstance(p, (str, dict))]
    if name.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(text))
        return [{(k or "").strip().lower(): (v or "") for k, v in r.items()} for r in reader]
    return [{"keyword": line} for line in text.splitlines()]   # .txt / manual list: one per line


def import_rows(db: Session, *, organization_id: UUID, ko_id: UUID, rows: list[dict], source: str,
                batch: str | None = None) -> dict:
    _ko(db, organization_id, ko_id)
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=422, detail=f"At most {MAX_IMPORT_ROWS} rows per import")
    batch = batch or uuid.uuid4().hex[:12]
    inserted, duplicates, rejected = 0, 0, []
    for n, row in enumerate(rows, start=1):
        kw = next((str(row[c]).strip() for c in KEYWORD_COLUMNS if row.get(c) not in (None, "")), "")
        if not kw:
            rejected.append({"row": n, "reason": "no keyword/query value"}); continue
        if len(kw) > 500:
            rejected.append({"row": n, "reason": "longer than 500 characters"}); continue
        kind = normalize(str(row.get("type") or row.get("kind") or "primary")).replace(" ", "_").replace("-", "_")
        if kind not in KEYWORD_KINDS:
            rejected.append({"row": n, "reason": f"unknown type '{kind}'"}); continue
        stmt = insert(KnowledgeKeyword).values(
            id=uuid.uuid4(), organization_id=organization_id, knowledge_object_id=ko_id, keyword=kw, normalized=normalize(kw),
            kind=kind, intent=(row.get("intent") or None), source=(row.get("source") or source), volume=_int_or_none(row.get("volume")),
            country=(row.get("country") or None), language=(row.get("language") or None), notes=(row.get("notes") or None),
            priority=(row.get("priority") or None), import_batch=batch, raw={k: (str(v) if v is not None else None) for k, v in row.items()},
            research_task_ids=[],
        ).on_conflict_do_nothing(constraint="uq_knowledge_keywords_dedupe").returning(KnowledgeKeyword.id)
        # RETURNING, not rowcount: the driver reports -1 for this statement form.
        if db.execute(stmt).scalar_one_or_none() is not None:
            inserted += 1
        else:
            duplicates += 1
    db.commit()
    return {"batch": batch, "received": len(rows), "inserted": inserted, "duplicates": duplicates, "rejected": rejected}


def search_new_sources(db: Session, *, organization_id: UUID, ko_id: UUID, keyword_ids: list[UUID],
                       actor_user_id: UUID, research_project_id: UUID | None = None) -> dict:
    """Keyword -> query -> research task (source discovery). Traceability is
    stored on the keyword (research_task_ids). Execution uses the configured
    research provider; with none configured the task reports NOT_CONFIGURED."""
    ko = _ko(db, organization_id, ko_id)
    kws = db.execute(select(KnowledgeKeyword).where(KnowledgeKeyword.organization_id == organization_id,
                                                    KnowledgeKeyword.knowledge_object_id == ko_id,
                                                    KnowledgeKeyword.id.in_(keyword_ids))).scalars().all()
    if len(kws) != len(set(keyword_ids)):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Keyword not found")
    if research_project_id is None:
        project = research_project_service.create_research_project(
            db, organization_id=organization_id, name=f"Keyword source discovery: {ko.canonical_name}"[:300],
            objective=f"Find new sources for {ko.canonical_name} from selected keywords.", mode=ResearchMode.USER_DRIVEN,
            knowledge_gap_id=None, actor_user_id=actor_user_id)
        research_project_id = project.id
    created = []
    for kw in kws:
        task = research_task_service.create_research_task(
            db, organization_id=organization_id, research_project_id=research_project_id, task_type="source_discovery",
            query=f"{kw.keyword} {ko.canonical_name}" if normalize(ko.canonical_name) not in kw.normalized else kw.keyword,
            actor_user_id=actor_user_id)
        kw.research_task_ids = list(kw.research_task_ids or []) + [str(task.id)]
        # REQ-KW-004: first link of the canonical traceability chain.
        from app.models.keyword_trace import STAGE_RESEARCH, KeywordTrace
        db.add(KeywordTrace(organization_id=organization_id, keyword_id=kw.id, stage=STAGE_RESEARCH,
                            research_task_id=task.id))
        created.append({"keyword_id": str(kw.id), "research_task_id": str(task.id), "query": task.query})
    db.commit()
    return {"research_project_id": str(research_project_id), "tasks": created}
